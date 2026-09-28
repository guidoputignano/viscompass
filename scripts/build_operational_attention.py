"""Build reference evidence only; never publish organization facts or spending.

Run with bundled Python and --local-workbook to include source-reported PHT.
Raw downloads are immutable inputs, outside public/. Snapshot dates are explicit.
"""
import argparse, csv, hashlib, io, json, re
from pathlib import Path
from collections import defaultdict

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data/raw/operational-attention'
SOURCES = {
 'h': ('class-h-20260430.csv','2026-04-30','https://www.aifa.gov.it/documents/20142/3815901/Classe_H_per_principio_attivo_30-04-2026.csv'),
 'a': ('class-a-20260430.csv','2026-04-30','https://www.aifa.gov.it/documents/20142/3815901/Classe_A_per_principio_attivo_30-04-2026.csv'),
 'registry': ('registries-20260923.csv','2026-09-23','https://www.aifa.gov.it/documents/20142/3258518/Elenco_Registri_PT_attivi_23.09.2026.csv'),
 'shortage': ('shortages-20260925.csv','2026-09-25','https://www.aifa.gov.it/documents/20142/847339/elenco_medicinali_carenti.csv'),
}

def aic(value):
 value=str(value).strip()
 if not re.fullmatch(r'\d{1,9}',value): raise ValueError('Invalid source AIC: '+value)
 return value.zfill(9)

def main():
 parser=argparse.ArgumentParser();parser.add_argument('--local-workbook');args=parser.parse_args()
 tables={};meta={}
 for key,(file,date,url) in SOURCES.items():
  raw=(RAW/file).read_bytes()
  try: text=raw.decode('utf-8-sig')
  except UnicodeDecodeError: text=raw.decode('cp1252')
  rows=list(csv.reader(io.StringIO(text),delimiter=';'))
  if key=='shortage':
   assert '25/09/2026' in text[:1000], 'Refresh source date before rebuilding'
   rows=rows[2:]
  header=[' '.join(x.split()) for x in rows[0]]
  tables[key]=[dict(zip(header,r)) for r in rows[1:] if any(r)]
  assert len(tables[key])>100, 'Unexpected/truncated source '+key
  meta[key]={'date':date,'url':url,'sha256':hashlib.sha256(raw).hexdigest(),'rows':len(tables[key])}
 h=sorted({aic(r['Codice AIC']) for r in tables['h']});a=sorted({aic(r['AIC']) for r in tables['a']})
 shortage={}
 for r in tables['shortage']:
  key=aic(r['Codice AIC'])
  shortage.setdefault(key,[]).append({'start':r['Data inizio'],'reason':r['Motivazioni']})
 registry=[]
 for r in tables['registry']:
  if r['Stato del Monitoraggio'].strip().lower()!='attivo':continue
  registry.append({'atc':r['Codice ATC'].strip().upper(),'name':r['Farmaco'].strip().upper(),
   'indication':r['Indicazione rimborsata'],'kind':r['Tipologia di Monitoraggio'],'url':r['Scheda']})
 pht={}
 if args.local_workbook:
  from openpyxl import load_workbook
  file=Path(args.local_workbook);w=load_workbook(file,read_only=True,data_only=True)
  sheet=w['DB TOT'];it=sheet.iter_rows(values_only=True);header=next(it)
  assert header[4]=='COD AIC' and header[23]=='Flag PHT (1)', 'Local schema changed'
  evidence=defaultdict(set)
  for r in it:
   # Latest local reference period only; no historical claim and no org data retained.
   if r[0]!=2026 or not str(r[17]).startswith('J01'):continue
   if not re.fullmatch(r'\d{1,9}',str(r[4]).strip()):continue
   evidence[aic(r[4])].add(str(r[23] or '').strip())
  pht={k:('yes' if v=={'S'} else 'no' if v=={'N'} else 'unknown') for k,v in evidence.items()}
  meta['pht']={'date':'2026-05','label':'Flag PHT nel file locale, gennaio–maggio 2026',
   'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':len(pht)}
 out={'sources':meta,'h':h,'a':a,'shortage':shortage,'registry':registry,'pht':pht}
 dest=ROOT/'data/derived/operational-attention.json';dest.parent.mkdir(parents=True,exist_ok=True)
 dest.write_text(json.dumps(out,ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8')
 print(json.dumps({'sources':meta,'pht_statuses':{s:list(pht.values()).count(s) for s in ['yes','no','unknown']}}))

if __name__=='__main__':main()
